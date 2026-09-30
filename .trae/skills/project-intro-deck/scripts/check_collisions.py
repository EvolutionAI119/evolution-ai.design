# -*- coding: utf-8 -*-
"""PPT 碰撞检测：用法  python check_collisions.py <file.pptx>
检测：形状越出幻灯片边界、文本框两两相交、图片两两相交。
注意：文本框为无填充透明占位，相交即视为潜在字符重叠风险。"""
import sys
from pptx import Presentation

if len(sys.argv) != 2:
    print("usage: python check_collisions.py <file.pptx>")
    sys.exit(2)

prs = Presentation(sys.argv[1])
SW, SH = prs.slide_width, prs.slide_height
TOL = 5000  # EMU，约 0.005 英寸，容许描边误差


def box(sh):
    return sh.left, sh.top, sh.left + sh.width, sh.top + sh.height


def inter(a, b):
    return not (a[2] <= b[0] or b[2] <= a[0] or a[3] <= b[1] or b[3] <= a[1])


issues = 0
for i, slide in enumerate(prs.slides, 1):
    shapes = [sh for sh in slide.shapes if sh.left is not None]
    tboxes = [sh for sh in shapes
              if sh.has_text_frame and sh.text_frame.text.strip()]
    pics = [sh for sh in shapes if sh.shape_type == 13]

    for sh in shapes:
        x1, y1, x2, y2 = box(sh)
        if x1 < -TOL or y1 < -TOL or x2 > SW + TOL or y2 > SH + TOL:
            label = sh.text_frame.text[:25] if sh.has_text_frame else "(shape)"
            print("P%d OFFSLIDE: %s" % (i, label.replace("\n", " ")))
            issues += 1

    for ai in range(len(tboxes)):
        for bi in range(ai + 1, len(tboxes)):
            if inter(box(tboxes[ai]), box(tboxes[bi])):
                ta = tboxes[ai].text_frame.text[:20].replace("\n", " ")
                tb = tboxes[bi].text_frame.text[:20].replace("\n", " ")
                print("P%d TEXT-OVERLAP: [%s] x [%s]" % (i, ta, tb))
                issues += 1

    for ai in range(len(pics)):
        for bi in range(ai + 1, len(pics)):
            if inter(box(pics[ai]), box(pics[bi])):
                print("P%d PIC-OVERLAP" % i)
                issues += 1

print("issues:", issues)
sys.exit(1 if issues else 0)
